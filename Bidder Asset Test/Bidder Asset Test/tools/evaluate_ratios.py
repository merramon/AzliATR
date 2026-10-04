"""Compute Current and Quick ratios per bidder and write the Excel evaluation.

Usage:
    .venv/bin/python tools/evaluate_ratios.py [--input .tmp/financials] [--overrides inputs/overrides.csv]
        [--current-min 1.0] [--quick-min 1.0] [--out outputs/asset_test_results.xlsx]

Current ratio = Current Assets / Current Liabilities
Quick ratio   = (Current Assets - Inventory) / Current Liabilities

Overrides CSV (optional) columns: bidder,year,field,value[,reason]
  field is one of current_assets, current_liabilities, inventory. Overrides win over parsed values
  and count as high confidence (a human checked them).

Result per bidder (latest year): PASS, FAIL, or REVIEW (missing / low-confidence figures).
"""
import argparse
import csv
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

FIELDS = ["current_assets", "current_liabilities", "inventory"]
FILLS = {
    "PASS": PatternFill("solid", fgColor="C6EFCE"),
    "FAIL": PatternFill("solid", fgColor="FFC7CE"),
    "REVIEW": PatternFill("solid", fgColor="FFEB9C"),
}


def load_overrides(path: Path) -> dict:
    ov = {}
    if not path.exists():
        return ov
    with path.open(newline="") as f:
        for row in csv.DictReader(f):
            if not row.get("bidder") or not row.get("value"):
                continue
            key = (row["bidder"].strip(), row["year"].strip(), row["field"].strip())
            ov[key] = {"value": float(row["value"].replace(",", "")), "reason": (row.get("reason") or "").strip()}
    return ov


def evaluate(fin: dict, overrides: dict, cur_min: float, quick_min: float) -> list[dict]:
    bidder = fin["bidder"]
    years = set(fin.get("years") or [])
    for f in fin["fields"].values():
        years.update(f["values"].keys())
    years.update(y for (b, y, _) in overrides if b == bidder)
    rows = []
    for year in sorted(years, reverse=True):
        row = {"bidder": bidder, "year": year, "notes": [], "low": False}
        for field in FIELDS:
            ov = overrides.get((bidder, year, field))
            parsed = fin["fields"].get(field)
            if ov:
                row[field] = ov["value"]
                row[f"{field}_src"] = f"override: {ov['reason'] or 'manual'}"
            elif parsed and year in parsed["values"]:
                row[field] = parsed["values"][year]
                row[f"{field}_src"] = f"p.{parsed['page']}: {parsed['line']}"
                if parsed["confidence"] == "low":
                    row["low"] = True
                    row["notes"].append(f"{field} low confidence ({parsed['note']})")
            else:
                row[field] = None
                row[f"{field}_src"] = ""
        ca, cl, inv = row["current_assets"], row["current_liabilities"], row["inventory"]
        if inv is None and ca is not None:
            inv = 0.0
            row["notes"].append("inventory not found - assumed 0")
        if ca is None or cl is None:
            row["current_ratio"] = row["quick_ratio"] = None
            row["result"] = "REVIEW"
            row["notes"].append("missing " + ", ".join(n for n, v in [("current assets", ca), ("current liabilities", cl)] if v is None))
        elif cl <= 0:
            row["current_ratio"] = row["quick_ratio"] = None
            row["result"] = "REVIEW"
            row["notes"].append("current liabilities is zero/negative - check figure")
        else:
            row["current_ratio"] = ca / cl
            row["quick_ratio"] = (ca - inv) / cl
            passed = row["current_ratio"] >= cur_min and row["quick_ratio"] >= quick_min
            row["result"] = "REVIEW" if row["low"] else ("PASS" if passed else "FAIL")
            if row["low"]:
                row["notes"].append(f"would be {'PASS' if passed else 'FAIL'} if figures confirmed")
        rows.append(row)
    if not rows:
        rows.append({"bidder": bidder, "year": "", "current_assets": None, "current_liabilities": None,
                     "inventory": None, "current_ratio": None, "quick_ratio": None, "result": "REVIEW",
                     "notes": list(fin.get("issues", [])) or ["no figures found"]})
    else:
        latest = rows[0]
        fully_overridden = all(latest.get(f"{f}_src", "").startswith("override") for f in FIELDS)
        if not fully_overridden:
            latest["notes"] = [i for i in fin.get("issues", []) if "inventory" not in i] + latest["notes"]
    return rows


def member_tag(bidder: str) -> str:
    return bidder.replace("Bidder ", "")  # "Bidder A1" -> "A1"


def combine_jv(label: str, members: list[list[dict]], cur_min: float, quick_min: float) -> dict:
    """Joint venture: one bid, several companies. Ratios use the members' latest-year figures added together."""
    latest = [rows[0] for rows in members]
    tags = [member_tag(m["bidder"]) for m in latest]
    row = {"bidder": label, "jv": True, "member_rows": members,
           "company": " + ".join(m.get("company") or member_tag(m["bidder"]) for m in latest),
           "notes": [f"Joint venture - combined figures of {' + '.join(tags)}"]}
    years = [m["year"] for m in latest]
    row["year"] = years[0] if len(set(years)) == 1 else " / ".join(f"{t}: {y}" for t, y in zip(tags, years))
    if len(set(years)) > 1:
        row["notes"].append("members' latest financial years differ - each member's latest year is used")
    missing = [t for t, m in zip(tags, latest) if m["current_assets"] is None or m["current_liabilities"] is None]
    if missing:
        row.update(current_assets=None, current_liabilities=None, inventory=None, current_ratio=None,
                   quick_ratio=None, result="REVIEW")
        row["notes"].append(f"figures missing for {', '.join(missing)}")
        return row
    ca = sum(m["current_assets"] for m in latest)
    cl = sum(m["current_liabilities"] for m in latest)
    inv = sum(m["inventory"] or 0.0 for m in latest)
    row.update(current_assets=ca, current_liabilities=cl, inventory=inv)
    if cl <= 0:
        row.update(current_ratio=None, quick_ratio=None, result="REVIEW")
        row["notes"].append("combined current liabilities is zero/negative - check figures")
        return row
    row["current_ratio"], row["quick_ratio"] = ca / cl, (ca - inv) / cl
    passed = row["current_ratio"] >= cur_min and row["quick_ratio"] >= quick_min
    unsure = [t for t, m in zip(tags, latest) if m["result"] == "REVIEW"]
    if unsure:
        row["result"] = "REVIEW"
        row["notes"].append(f"{', '.join(unsure)} needs checking - would be {'PASS' if passed else 'FAIL'} if figures confirmed")
    else:
        row["result"] = "PASS" if passed else "FAIL"
    return row


def group_jvs(all_rows: list[list[dict]], groups: dict, cur_min: float, quick_min: float) -> list[list[dict]]:
    """Replace JV members by one combined entry, placed where the first member was."""
    by_bidder = {rows[0]["bidder"]: rows for rows in all_rows}
    member_of = {m: label for label, ms in groups.items() for m in ms}
    out, done = [], set()
    for rows in all_rows:
        label = member_of.get(rows[0]["bidder"])
        if not label:
            out.append(rows)
        elif label not in done:
            done.add(label)
            members = [by_bidder[m] for m in groups[label] if m in by_bidder]
            out.append([combine_jv(label, members, cur_min, quick_min)])
    return out


def display_name(r: dict) -> str:
    return f"{r['bidder']} (JV)" if r.get("jv") else r["bidder"]


def write_excel(all_rows: list[list[dict]], out: Path, cur_min: float, quick_min: float):
    wb = Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Summary"
    ws.append([f"Asset test evaluation - {date.today().isoformat()}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Pass criteria: Current ratio >= {cur_min}  AND  Quick ratio >= {quick_min}  (latest year)"])
    ws.append([])
    hdr = ["Bidder", "Company", "Year", "Current Assets", "Current Liabilities", "Inventory",
           "Current Ratio", "Quick Ratio", "Result", "Notes"]
    ws.append(hdr)
    for c in ws[4]:
        c.font = bold
    for rows in all_rows:
        r = rows[0]  # latest year
        ws.append([display_name(r), r.get("company", ""), r["year"], r["current_assets"], r["current_liabilities"], r["inventory"],
                   r["current_ratio"], r["quick_ratio"], r["result"], "; ".join(r["notes"])])
        ws.cell(ws.max_row, 9).fill = FILLS[r["result"]]
        for m in [mr[0] for mr in r.get("member_rows", [])]:
            ws.append([f"    {member_tag(m['bidder'])}", m.get("company") or "", m["year"], m["current_assets"],
                       m["current_liabilities"], m["inventory"], m["current_ratio"], m["quick_ratio"],
                       f"({m['result']})", "JV member - for information; the bid is judged on the combined figures. "
                       + "; ".join(m["notes"])])
            ws.cell(ws.max_row, 1).font = Font(italic=True)
            ws.cell(ws.max_row, 2).font = Font(italic=True)

    wd = wb.create_sheet("All Years")
    wd.append(hdr + ["Source: Current Assets", "Source: Current Liabilities", "Source: Inventory"])
    for c in wd[1]:
        c.font = bold
    flat = []
    for rows in all_rows:
        for mr in rows[0].get("member_rows", []):
            flat.append(mr)
        flat.append(rows)
    for rows in flat:
        for r in rows:
            wd.append([display_name(r), r.get("company", ""), r["year"], r["current_assets"], r["current_liabilities"], r["inventory"],
                       r["current_ratio"], r["quick_ratio"], r["result"], "; ".join(r["notes"]),
                       r.get("current_assets_src", ""), r.get("current_liabilities_src", ""), r.get("inventory_src", "")])
            wd.cell(wd.max_row, 9).fill = FILLS[r["result"]]

    for sheet, first in ((ws, 5), (wd, 2)):
        for row in sheet.iter_rows(min_row=first):
            for c in row[3:6]:
                c.number_format = "#,##0.00"
            for c in row[6:8]:
                c.number_format = "0.00"
        for col, width in zip("ABCDEFGHIJKLM", [16, 40, 8, 16, 18, 14, 13, 12, 10, 60, 50, 50, 50]):
            sheet.column_dimensions[col].width = width
        for row in sheet.iter_rows(min_row=first):
            for c in row[9:]:
                c.alignment = Alignment(wrap_text=True, vertical="top")

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", default=".tmp/financials")
    ap.add_argument("--overrides", default="inputs/overrides.csv")
    ap.add_argument("--current-min", type=float, help="default: config/asset_test.json")
    ap.add_argument("--quick-min", type=float, help="default: config/asset_test.json")
    ap.add_argument("--out", default="outputs/asset_test_results.xlsx")
    ap.add_argument("--json", help="also write the results as JSON to this path (used by the web app)")
    ap.add_argument("--names", help='optional JSON {"Bidder A": "Company Sdn Bhd"} - shown next to the bidder label')
    ap.add_argument("--groups", help='optional JSON of joint ventures {"Bidder A": ["Bidder A1", "Bidder A2"]} - '
                                     'members are combined into one bid')
    args = ap.parse_args()
    cfg_path = Path("config/asset_test.json")
    cfg = json.loads(cfg_path.read_text()) if cfg_path.exists() else {}
    if args.current_min is None:
        args.current_min = (cfg.get("current_ratio") or {}).get("minimum")
    if args.quick_min is None:
        args.quick_min = (cfg.get("quick_ratio") or {}).get("minimum")
    if args.current_min is None or args.quick_min is None:
        print("Pass marks not set - fill in the minimums in config/asset_test.json or pass --current-min/--quick-min",
              file=sys.stderr)
        return 1

    files = sorted(Path(args.input).glob("*.json"))
    if not files:
        print(f"No parsed financials in {args.input} - run parse_financials.py first", file=sys.stderr)
        return 1
    overrides = load_overrides(Path(args.overrides))
    all_rows = [evaluate(json.loads(f.read_text()), overrides, args.current_min, args.quick_min) for f in files]
    names = json.loads(Path(args.names).read_text()) if args.names else {}
    groups = json.loads(Path(args.groups).read_text()) if args.groups else {}
    jv_members = {m for ms in groups.values() for m in ms}
    detected = {json.loads(f.read_text())["bidder"]: json.loads(f.read_text()).get("company_name", "") for f in files}
    for rows in all_rows:
        for r in rows:
            # JV members' names are read from their reports unless given
            # company name read from the report, unless the user typed one
            r["company"] = names.get(r["bidder"]) or detected.get(r["bidder"], "")
    all_rows = group_jvs(all_rows, groups, args.current_min, args.quick_min)
    write_excel(all_rows, Path(args.out), args.current_min, args.quick_min)
    def clean(r):
        r = {k: v for k, v in r.items() if k != "low"}
        if "member_rows" in r:
            r["member_rows"] = [[clean(x) for x in rows] for rows in r["member_rows"]]
        return r

    if args.json:
        Path(args.json).write_text(json.dumps({
            "current_min": args.current_min, "quick_min": args.quick_min,
            "bidders": [[clean(r) for r in rows] for rows in all_rows],
        }, indent=2))

    for rows in all_rows:
        r = rows[0]
        cr = f"{r['current_ratio']:.2f}" if r["current_ratio"] is not None else "-"
        qr = f"{r['quick_ratio']:.2f}" if r["quick_ratio"] is not None else "-"
        print(f"{r['result']:6} {r['bidder']} ({r['year'] or '?'}): current {cr}, quick {qr}")
    print(f"\nSaved {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
